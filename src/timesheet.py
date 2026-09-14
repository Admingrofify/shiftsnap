"""ShiftSnap timesheet generator: builds a clean, employer-ready Excel timesheet.

Reads shifts from ShiftStore and writes a formatted .xlsx with per-day rows
and a totals row. Worker details come from profile.json (see profile.sample.json).
"""
from __future__ import annotations

import datetime
import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .store import ShiftStore, fmt_duration, net_minutes, normalize_date, to_minutes

BASE = Path(__file__).resolve().parent.parent
PROFILE = BASE / "profile.json"
PROFILE_SAMPLE = BASE / "profile.sample.json"
OUT_DIR = BASE / "timesheets"

HEADERS = ["Date", "Day", "Time In", "Time Out", "Break",
           "Total Hours", "Hours (decimal)", "Approved (8h)", "Extra Hours", "Comment"]

HEADER_FILL = PatternFill("solid", fgColor="1F4E5F")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(bold=True, size=14, color="1F4E5F")
THIN = Side(style="thin", color="B0B0B0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")


def load_profile() -> dict:
    path = PROFILE if PROFILE.exists() else PROFILE_SAMPLE
    return json.loads(path.read_text())


def _write_sheet(ws, worker_name: str, shifts: dict, start: str, end: str,
                 profile: dict, hourly_rate=None) -> None:
    """Fill one worksheet with a worker's timesheet. Shared by solo & team exports."""
    ws.title = worker_name[:31]
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    ws.merge_cells("A1:J1")
    c = ws["A1"]
    c.value = f"TIMESHEET  —  {worker_name}  ({start} to {end})"
    c.font = TITLE_FONT
    c.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:J2")
    rate = hourly_rate if hourly_rate is not None else profile.get("hourly_rate", "")
    meta = (f"Classification: {profile.get('classification', '')}   |   "
            f"Hourly rate: ${rate}/hr   |   "
            f"Project: {profile.get('project_code', '')}   |   "
            f"Location: {profile.get('location', '')}")
    ws["A2"] = meta
    ws["A2"].alignment = LEFT
    ws.row_dimensions[2].height = 20

    for col, h in enumerate(HEADERS, start=1):
        cell = ws.cell(row=4, column=col, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = BORDER
    ws.row_dimensions[4].height = 22

    widths = [12, 10, 10, 10, 14, 12, 14, 12, 12, 30]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    row = 5
    total_min = 0
    approved_min_total = 0
    day = datetime.date.fromisoformat(start)
    end_d = datetime.date.fromisoformat(end)
    while day <= end_d:
        key = day.isoformat()
        for s in shifts.get(key, []):
            mins = net_minutes(s)
            total_min += mins
            approved = min(mins, 8 * 60)
            approved_min_total += approved
            brk = ""
            if s.get("break_start") and s.get("break_end"):
                brk = f"{s['break_start']}-{s['break_end']}"
            vals = [key, day.strftime("%a"), s["time_in"], s["time_out"] or "—", brk,
                    fmt_duration(mins), round(mins / 60, 2),
                    fmt_duration(approved), fmt_duration(mins - approved),
                    s.get("comment", "")]
            for col, v in enumerate(vals, start=1):
                cell = ws.cell(row=row, column=col, value=v)
                cell.border = BORDER
                cell.alignment = CENTER if col < 10 else LEFT
            if day.weekday() >= 5:
                for col in range(1, 11):
                    ws.cell(row=row, column=col).fill = PatternFill("solid", fgColor="F2F2F2")
            row += 1
        day += datetime.timedelta(days=1)

    for col, v in enumerate(
            ["", "", "", "", "TOTALS:", fmt_duration(total_min),
             round(total_min / 60, 2), fmt_duration(approved_min_total),
             fmt_duration(total_min - approved_min_total), ""], start=1):
        cell = ws.cell(row=row, column=col, value=v)
        cell.font = Font(bold=True)
        cell.border = BORDER
        cell.alignment = CENTER if col < 10 else LEFT
        cell.fill = PatternFill("solid", fgColor="D9E8EC")
    ws.row_dimensions[row].height = 22

    rate = hourly_rate if hourly_rate is not None else (profile.get("hourly_rate") or 0)
    est = round(total_min / 60 * rate, 2)
    ws.merge_cells(f"A{row + 2}:J{row + 2}")
    ws[f"A{row + 2}"] = f"Estimated gross pay: ${est:,.2f}  ({round(total_min / 60, 2)}h × ${rate}/h)"
    ws[f"A{row + 2}"].font = Font(bold=True, size=11)


def generate_timesheet(start: str, end: str, store: ShiftStore | None = None,
                     worker_name: str | None = None,
                     hourly_rate=None) -> str:
    """Generate the timesheet Excel for start..end (YYYY-MM-DD). Returns file path."""
    start = normalize_date(start)
    end = normalize_date(end)
    store = store or ShiftStore()
    profile = load_profile()
    shifts = store.list_shifts(start, end)

    wb = openpyxl.Workbook()
    _write_sheet(wb.active, worker_name or profile.get("name", "Timesheet"),
                 shifts, start, end, profile, hourly_rate)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    fname = f"Timesheet_{start}_to_{end}.xlsx"
    path = OUT_DIR / fname
    wb.save(path)
    return str(path)


def generate_team_timesheet(start: str, end: str, store) -> str:
    """Admin export: one workbook, a summary sheet plus one sheet per worker."""
    start = normalize_date(start)
    end = normalize_date(end)
    profile = load_profile()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Team Summary"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    headers = ["Worker", "Days", "Total Hours", "Rate ($/h)", "Est. Pay ($)",
               "Approved (8h/d)", "Extra Hours"]
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = BORDER
    summaries = store.team_summary(start, end)
    for r, s in enumerate(summaries, start=2):
        rate = s.get("hourly_rate") or 0
        hours = s["total_decimal"]
        pay = round(hours * rate, 2)
        for col, v in enumerate(
                [s["worker"], s["days_worked"], s["total_hours"], rate, pay,
                 s["approved_hours"], s["extra_hours"]], start=1):
            cell = ws.cell(row=r, column=col, value=v)
            cell.border = BORDER
            cell.alignment = CENTER
    for i, w in enumerate([18, 10, 14, 12, 14, 14, 14], start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    for s in summaries:
        worker = next(w for w in store.list_workers()
                      if w["name"] == s["worker"])
        shifts = store.list_shifts(start, end, employee_id=worker["id"])
        _write_sheet(wb.create_sheet(), s["worker"], shifts, start, end,
                     profile, s.get("hourly_rate"))

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    fname = f"Team_Timesheet_{start}_to_{end}.xlsx"
    path = OUT_DIR / fname
    wb.save(path)
    return str(path)
